import pandas as pd
from docx import Document

def replace_placeholder_in_paragraph(paragraph, placeholder, replacement_text):
    # Iterate through runs in the paragraph
    for run in paragraph.runs:
        if placeholder in run.text:
            # Replace placeholder with replacement text
            run.text = run.text.replace(placeholder, replacement_text)

def generate_documents_from_csv(csv_file, template_docx, output_folder):
    # Load the CSV data
    df = pd.read_csv(csv_file)
    
    # Iterate through each row in the DataFrame
    # This expects a csv where all the info is in each row
    for index, row in df.iterrows():
        # Load the template document
        doc = Document(template_docx)
        
        # Define the mapping of placeholders to replacement text
        placeholder_to_text = {
            'ORGNAME': row['org_name'],
            'MONTH': row['month'],
            'YEAR': str(row['year']),
            'TOTAL_ATTENDANCE':str(row['total_attendance'])
        }

        # Replace placeholders with text
        for para in doc.paragraphs:
            for placeholder, replacement_text in placeholder_to_text.items():
                replace_placeholder_in_paragraph(para, placeholder, replacement_text)

        # Save the modified document
        output_file = f"{output_folder}/{row['org_name']}_document.docx"
        doc.save(output_file)
        print(f"Saved: {output_file}")

# Example usage
generate_documents_from_csv('total_attendances.csv', 'template_document.docx', 'output_folder')
